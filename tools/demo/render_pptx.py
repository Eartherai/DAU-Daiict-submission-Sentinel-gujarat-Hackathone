"""Build a 16:9 PowerPoint from the rendered deck PNGs.

The portal accepts PPT or PDF. Each slide is a full-bleed 1920×1080 page so
the live UI fills the screen in slideshow, not a letterboxed thumbnail.

    python tools/demo/render_pptx.py
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from saakshya.common.paths import display


def main() -> int:
    from pptx import Presentation
    from pptx.util import Emu, Inches

    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--src", default="var/demo/SAAKSHYA_deck")
    ap.add_argument("--out", default="var/demo/SAAKSHYA_deck.pptx")
    a = ap.parse_args()
    src = Path(a.src)
    slides = sorted(src.glob("slide_*.png"))
    if not slides:
        print(f"no slides in {src}", file=sys.stderr)
        return 2

    # Widescreen 16:9 at 1920×1080 px (914400 EMU per inch).
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    w, h = prs.slide_width, prs.slide_height
    for png in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), Emu(0), Emu(0), width=w, height=h)
        print(f"  {png.name}")

    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"\npptx : {display(out)} ({len(slides)} slides, {out.stat().st_size/1e6:.1f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
